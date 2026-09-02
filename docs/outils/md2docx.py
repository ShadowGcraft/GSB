#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Conversion du cahier des charges GSB, du Markdown vers un document Word.

Le document source contient des elements que les convertisseurs generiques rendent mal :
des tableaux dans des citations, des blocs Mermaid, une numerotation RG/F/A a preserver.
Ce script traite chaque cas explicitement plutot que de passer par une conversion generique.
"""
import re
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else 'docs/cahier-des-charges.md'
DST = sys.argv[2] if len(sys.argv) > 2 else 'docs/cahier-des-charges.docx'
IMG = sys.argv[3] if len(sys.argv) > 3 else None  # dossier des diagrammes rendus

BLEU = RGBColor(0x1F, 0x3B, 0x63)
GRIS = RGBColor(0x59, 0x59, 0x59)
ROUGE = RGBColor(0x9B, 0x2C, 0x2C)


# --------------------------------------------------------------------------- styles

def bordure(par, cote='left', taille=18, couleur='1F3B63'):
    """Ajoute une bordure a un paragraphe. python-docx ne l'expose pas."""
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        pPr.append(bdr)
    el = OxmlElement('w:' + cote)
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(taille))
    el.set(qn('w:space'), '8')
    el.set(qn('w:color'), couleur)
    bdr.append(el)


def fond(par_ou_cell, couleur='F2F4F7'):
    """Applique une trame de fond a un paragraphe ou a une cellule."""
    el = par_ou_cell._p.get_or_add_pPr() if hasattr(par_ou_cell, '_p') else par_ou_cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), couleur)
    el.append(sh)


def fond_cellule(cell, couleur):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), couleur)
    tcPr.append(sh)


def prepare_styles(doc):
    n = doc.styles['Normal']
    n.font.name = 'Calibri'
    n.font.size = Pt(10.5)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15

    for niveau, taille, gras in ((1, 20, True), (2, 15, True), (3, 12.5, True), (4, 11, True)):
        s = doc.styles['Heading %d' % niveau]
        s.font.name = 'Calibri'
        s.font.size = Pt(taille)
        s.font.bold = gras
        s.font.color.rgb = BLEU
        s.paragraph_format.space_before = Pt(16 if niveau <= 2 else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True

    for nom in ('List Bullet', 'List Number'):
        s = doc.styles[nom]
        s.font.name = 'Calibri'
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(3)


def pied_de_page(doc):
    """Numerotation des pages, champ PAGE insere a la main."""
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Cahier des charges GSB — ').font.size = Pt(8)
    for texte, instr in (('', 'PAGE'), (' / ', None), ('', 'NUMPAGES')):
        if instr is None:
            r = p.add_run(texte)
            r.font.size = Pt(8)
            continue
        r = p.add_run()
        r.font.size = Pt(8)
        deb = OxmlElement('w:fldChar')
        deb.set(qn('w:fldCharType'), 'begin')
        txt = OxmlElement('w:instrText')
        txt.set(qn('xml:space'), 'preserve')
        txt.text = ' %s ' % instr
        fin = OxmlElement('w:fldChar')
        fin.set(qn('w:fldCharType'), 'end')
        r._r.append(deb)
        r._r.append(txt)
        r._r.append(fin)
    for r in p.runs:
        r.font.color.rgb = GRIS


# --------------------------------------------------- rendu du markdown en ligne

MOTIF = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))', re.S)


SAUT = re.compile(r'<br\s*/?>', re.I)


def ecrire_inline(par, texte, base_gras=False, base_taille=None, couleur=None):
    """Ecrit du texte markdown en gerant gras, italique, code, liens et sauts de ligne HTML.

    Le document source emploie <br> a l'interieur des cellules de tableau, notation valide
    sur GitHub. Rendue telle quelle dans Word, la balise resterait visible.
    """
    texte = texte.replace('\n', ' ')

    fragments = SAUT.split(texte)
    if len(fragments) > 1:
        for n, frag in enumerate(fragments):
            if n:
                par.add_run().add_break()
            frag = frag.strip()
            if frag:
                ecrire_inline(par, frag, base_gras, base_taille, couleur)
        return

    for morceau in MOTIF.split(texte):
        if not morceau:
            continue
        gras, ital, code = base_gras, False, False
        contenu = morceau
        if morceau.startswith('***') and morceau.endswith('***'):
            contenu, gras, ital = morceau[3:-3], True, True
        elif morceau.startswith('**') and morceau.endswith('**'):
            contenu, gras = morceau[2:-2], True
        elif morceau.startswith('*') and morceau.endswith('*') and len(morceau) > 2:
            contenu, ital = morceau[1:-1], True
        elif morceau.startswith('`') and morceau.endswith('`'):
            contenu, code = morceau[1:-1], True
        else:
            lien = re.match(r'\[([^\]]+)\]\(([^)]+)\)', morceau)
            if lien:
                contenu = lien.group(1)
        r = par.add_run(contenu)
        r.bold = gras
        r.italic = ital
        if code:
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x2D, 0x5C)
        if base_taille:
            r.font.size = Pt(base_taille)
        if couleur is not None and not code:
            r.font.color.rgb = couleur
        # les marqueurs A ARBITRER et HYPOTHESE ressortent en rouge
        if gras and ('ARBITRER' in contenu or 'HYPOTH' in contenu.upper()):
            r.font.color.rgb = ROUGE


def cellules(ligne):
    l = ligne.strip()
    if l.startswith('> '):
        l = l[2:].strip()
    return [c.strip() for c in l.strip('|').split('|')]


def est_separateur(ligne):
    l = ligne.strip()
    if l.startswith('> '):
        l = l[2:].strip()
    return bool(re.match(r'^\|[\s:|-]+\|$', l))


# --------------------------------------------------------------- construction

def largeurs_colonnes(entete, corps, total_cm=16.0):
    """Repartit la largeur en fonction du contenu reel de chaque colonne.

    Word repartit les colonnes a egalite, ce qui casse un tableau ou une colonne porte
    une phrase et les autres un mot. On pondere par la longueur mediane du contenu,
    ecretee pour qu'une cellule exceptionnellement longue ne prenne pas tout.
    """
    n = len(entete)
    car = 0.19          # largeur moyenne d'un caractere a 9,5 points, en cm
    marge = 0.35        # padding interne d'une cellule Word

    poids, minima = [], []
    for i in range(n):
        textes = [SAUT.sub(' ', c[i]) for c in corps if i < len(c)] + [entete[i]]
        textes = [re.sub(r'[*`\[\]]', '', t) for t in textes]

        longueurs = sorted(len(t) for t in textes)
        median = longueurs[len(longueurs) // 2] if longueurs else 1
        poids.append(max(6.0, min(float(median), 70.0)) ** 0.72)

        # un mot ne doit jamais etre coupe en deux : il fixe la largeur plancher
        mots = [m for t in textes for m in t.split()]
        plus_long = max((len(m) for m in mots), default=3)
        minima.append(min(max(1.5, plus_long * car + marge), 4.2))

    somme = sum(poids) or 1.0
    larg = [total_cm * p / somme for p in poids]

    # on redistribue jusqu'a ce que chaque colonne atteigne son plancher
    for _ in range(6):
        manque = sum(minima[i] - larg[i] for i in range(n) if larg[i] < minima[i])
        if manque <= 0.01:
            break
        gros = [i for i in range(n) if larg[i] > minima[i] + 0.1]
        if not gros:
            break
        surplus = sum(larg[i] - minima[i] for i in gros) or 1.0
        for i in gros:
            larg[i] -= manque * (larg[i] - minima[i]) / surplus
        larg = [max(larg[i], minima[i]) for i in range(n)]

    # si la somme deborde malgre tout, on rabote proportionnellement
    if sum(larg) > total_cm:
        f = total_cm / sum(larg)
        larg = [x * f for x in larg]
    return larg


def ajoute_tableau(doc, lignes, indente=False):
    entete = cellules(lignes[0])
    corps = [cellules(l) for l in lignes[2:]]
    t = doc.add_table(rows=1, cols=len(entete))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    larg = largeurs_colonnes(entete, corps, 16.0 - (0.5 if indente else 0.0))
    for i, txt in enumerate(entete):
        c = t.rows[0].cells[i]
        c.text = ''
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        ecrire_inline(p, txt, base_gras=True, base_taille=9.5)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.bold = True
        fond_cellule(c, '1F3B63')
    for n, ligne in enumerate(corps):
        cells = t.add_row().cells
        for i, txt in enumerate(ligne[:len(entete)]):
            c = cells[i]
            c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            ecrire_inline(p, txt, base_taille=9.5)
            if n % 2 == 1:
                fond_cellule(c, 'F4F6F9')
    # une ligne ne se coupe pas entre deux pages, sinon une cellule longue est tronquee
    # en plein mot ; et la ligne d'entete se repete en haut de chaque page du tableau
    for n, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
        if n == 0:
            th = OxmlElement('w:tblHeader')
            trPr.append(th)

    # la largeur doit etre posee sur chaque cellule, Word ignore celle de la colonne seule
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if i < len(larg):
                c.width = Cm(larg[i])
    for i, col in enumerate(t.columns):
        if i < len(larg):
            col.width = Cm(larg[i])
    if indente:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.left_indent = Cm(0.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def ajoute_code(doc, lignes, titre=None):
    if titre:
        p = doc.add_paragraph()
        r = p.add_run(titre)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS
        p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    fond(p, 'F5F5F5')
    bordure(p, 'left', 18, 'B0B7C3')
    for i, l in enumerate(lignes):
        r = p.add_run(l)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        if i < len(lignes) - 1:
            r.add_break()
    return p


LARGEUR_MAX = 16.0   # cm, largeur utile de la page moins les marges
HAUTEUR_MAX = 19.5   # cm, pour laisser place a la legende sur la meme page


def taille_image(chemin):
    """Retourne la taille d'insertion en centimetres, bornee en largeur et en hauteur."""
    from PIL import Image
    px_l, px_h = Image.open(chemin).size
    larg = LARGEUR_MAX
    haut = larg * px_h / px_l
    if haut > HAUTEUR_MAX:
        haut = HAUTEUR_MAX
        larg = haut * px_l / px_h
    return larg, haut


def convertir(src, dst, dossier_images=None):
    texte = open(src, encoding='utf-8').read()
    lignes = texte.split('\n')

    doc = Document()
    prepare_styles(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    pied_de_page(doc)

    # index des diagrammes rendus, dans l'ordre d'apparition
    images = []
    if dossier_images and os.path.isdir(dossier_images):
        images = sorted(
            os.path.join(dossier_images, f)
            for f in os.listdir(dossier_images) if f.endswith('.png')
        )
    n_diag = 0
    legendes = {
        0: 'Figure 1 — Cas d’utilisation et acteurs',
        1: 'Figure 2 — Cycle de vie d’une fiche de frais',
        2: 'Figure 3 — Architecture applicative',
    }

    i = 0
    dans_entete = True
    while i < len(lignes):
        l = lignes[i]
        nu = l.strip()

        # --- page de garde : les metadonnees en tete du fichier
        if dans_entete and nu.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(90)
            r = p.add_run(nu[2:])
            r.bold = True
            r.font.size = Pt(30)
            r.font.color.rgb = BLEU
            i += 1
            continue

        if dans_entete and nu.startswith('**') and ' : ' in nu:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            ecrire_inline(p, nu, base_taille=11)
            i += 1
            continue

        if dans_entete and nu == '---':
            dans_entete = False
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        # --- blocs de code, y compris Mermaid
        if nu.startswith('```'):
            langue = nu[3:].strip()
            j = i + 1
            bloc = []
            while j < len(lignes) and not lignes[j].strip().startswith('```'):
                bloc.append(lignes[j])
                j += 1
            if langue == 'mermaid':
                if n_diag < len(images):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    # le diagramme est dimensionne par sa contrainte la plus forte,
                    # sans quoi un schema tres haut deborde de la page
                    larg, haut = taille_image(images[n_diag])
                    p.add_run().add_picture(images[n_diag], width=Cm(larg), height=Cm(haut))
                    lg = doc.add_paragraph()
                    lg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = lg.add_run(legendes.get(n_diag, 'Figure %d' % (n_diag + 1)))
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = GRIS
                    lg.paragraph_format.space_after = Pt(10)
                else:
                    ajoute_code(doc, bloc, 'Diagramme, source Mermaid :')
                n_diag += 1
            else:
                ajoute_code(doc, bloc)
            i = j + 1
            continue

        # --- separateur horizontal
        if nu == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            bordure(p, 'bottom', 6, 'C8CDD6')
            i += 1
            continue

        # --- titres
        m = re.match(r'^(#{1,6})\s+(.*)$', nu)
        if m:
            niveau = len(m.group(1))
            titre = m.group(2).strip()
            if niveau == 2 and re.match(r'^\d+\.', titre):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=min(niveau, 4))
            ecrire_inline(h, titre, base_gras=True)
            for r in h.runs:
                r.font.color.rgb = BLEU
            i += 1
            continue

        # --- tableaux, y compris a l'interieur d'une citation
        est_ligne = nu.startswith('|') or nu.startswith('> |')
        if est_ligne and i + 1 < len(lignes) and est_separateur(lignes[i + 1]):
            j = i
            bloc = []
            while j < len(lignes):
                s = lignes[j].strip()
                if not (s.startswith('|') or s.startswith('> |')):
                    break
                bloc.append(lignes[j])
                j += 1
            ajoute_tableau(doc, bloc, indente=nu.startswith('> '))
            i = j
            continue

        # --- citations : enonces des regles de gestion
        if nu.startswith('>'):
            j = i
            bloc = []
            while j < len(lignes) and lignes[j].strip().startswith('>'):
                if lignes[j].strip().startswith('> |') or est_separateur(lignes[j]):
                    break
                bloc.append(lignes[j].strip().lstrip('>').strip())
                j += 1
            if j == i:  # la citation commence par un tableau
                j = i + 1
            contenu = ' '.join(x for x in bloc if x).strip()
            if contenu:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)
                bordure(p, 'left', 14, '1F3B63')
                ecrire_inline(p, contenu)
            i = j
            continue

        # --- listes a puces
        m = re.match(r'^(\s*)[-*+]\s+(.*)$', l)
        if m:
            prof = len(m.group(1)) // 2
            p = doc.add_paragraph(style='List Bullet' if prof == 0 else 'List Bullet 2')
            ecrire_inline(p, m.group(2))
            i += 1
            continue

        # --- listes numerotees
        m = re.match(r'^(\s*)\d+[.)]\s+(.*)$', l)
        if m:
            p = doc.add_paragraph(style='List Number')
            ecrire_inline(p, m.group(2))
            i += 1
            continue

        # --- paragraphe : on recolle les lignes coupees a 100 colonnes
        if nu:
            bloc = []
            j = i
            while j < len(lignes):
                s = lignes[j].strip()
                if (not s or s.startswith(('#', '>', '|', '```', '---'))
                        or re.match(r'^\s*([-*+]|\d+[.)])\s', lignes[j])):
                    break
                bloc.append(s)
                j += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            ecrire_inline(p, ' '.join(bloc))
            i = j
            continue

        i += 1

    doc.save(dst)
    return doc, n_diag, len(images)


if __name__ == '__main__':
    d, nd, ni = convertir(SRC, DST, IMG)
    print('ecrit : %s' % DST)
    print('paragraphes : %d | tableaux : %d' % (len(d.paragraphs), len(d.tables)))
    print('diagrammes rencontres : %d | images inserees : %d' % (nd, min(nd, ni)))
